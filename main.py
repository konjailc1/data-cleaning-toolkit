"""Clean a customer-order CSV and write the result plus a short report.

Usage:
    python main.py sample_input/customers_messy.csv

Writes three files into output/ (or into --output-dir):
    cleaned.csv
    cleaned.xlsx
    report.md
"""

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cleaner import EXPECTED_COLUMNS, clean_orders


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Clean a messy customer-order CSV and write a report."
    )
    parser.add_argument("input_csv", help="Path to the messy CSV file")
    parser.add_argument(
        "--output-dir",
        default="output",
        help="Folder for cleaned.csv, cleaned.xlsx, and report.md (default: output)",
    )
    return parser.parse_args(argv)


def load_input(path: Path) -> pd.DataFrame:
    """Read every column as text.

    keep_default_na=False leaves tokens such as "n/a" as text, so the
    cleaner can count them instead of pandas hiding them as NaN.
    Phone numbers stay strings, so a leading 0 is not dropped.
    """
    if not path.is_file():
        raise FileNotFoundError(f"Input file not found: {path}")
    try:
        return pd.read_csv(path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
    except UnicodeDecodeError as error:
        raise UnicodeDecodeError(
            error.encoding,
            error.object,
            error.start,
            error.end,
            f"Could not read {path} as UTF-8. Save the CSV as UTF-8 and try again.",
        ) from error


def write_csv(df: pd.DataFrame, path: Path):
    # utf-8-sig lets Excel open Vietnamese names without mojibake.
    df.to_csv(path, index=False, encoding="utf-8-sig")


def format_worksheet(worksheet):
    """Bold header row, filter, frozen header, and fitted column widths."""
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="left", vertical="center")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    worksheet.row_dimensions[1].height = 22
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"

    text_headers = {"order_id", "phone"}
    for index, column_cells in enumerate(worksheet.columns, start=1):
        header = column_cells[0].value
        longest = len(str(header)) if header is not None else 0
        for cell in column_cells[1:]:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
            if header == "amount" and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            if header in text_headers:
                cell.number_format = "@"
        letter = get_column_letter(index)
        worksheet.column_dimensions[letter].width = min(longest + 3, 45)


def write_excel(df: pd.DataFrame, path: Path):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="cleaned")
        format_worksheet(writer.sheets["cleaned"])


def _markdown_cell(value) -> str:
    """Show a raw cell inside a markdown table, spaces included."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "(blank)"
    text = str(value).replace("\r", " ").replace("\n", " ").replace("|", "\\|")
    text = text.replace("`", "'")
    if text.strip() == "":
        return "`(blank)`"
    return f"`{text}`"


def build_report(stats, source_name: str) -> str:
    """Turn clean_orders stats into a markdown report."""
    invalid_rows = len({item["source_row"] for item in stats["invalid"]})
    lines = [
        "# Cleaning report",
        "",
        f"Source file: `{source_name}`",
        "",
        "Source row numbers below count the header as row 1, so the first",
        "data row is source row 2. That matches the line number in the CSV.",
        "",
        "| | |",
        "| --- | ---: |",
        f"| Rows before | {stats['rows_before']} |",
        f"| Rows after | {stats['rows_after']} |",
        f"| Duplicate rows removed | {stats['duplicates_removed']} |",
        f"| Invalid values cleared | {len(stats['invalid'])} |",
        f"| Rows with an invalid value | {invalid_rows} |",
        f"| Unrecognized cities kept for review | {len(stats['unmapped_cities'])} |",
        "",
        "Duplicate rows were removed only when every cleaned column matched.",
        "A row with a bad email or phone stays in the file; that cell is cleared.",
        "",
        "## Values fixed per column",
        "",
    ]

    if stats["fixes"]:
        lines.append("| Column | What changed | Count |")
        lines.append("| --- | --- | ---: |")
        for item in stats["fixes"]:
            lines.append(f"| {item['column']} | {item['action']} | {item['count']} |")
    else:
        lines.append("No values needed to change.")
    lines.append("")

    lines.extend(
        [
            "## Missing values",
            "",
            "Missing before is counted on the raw file. Missing after is counted",
            "on the cleaned file, after duplicates are removed. The after count",
            "also includes invalid emails, phones, dates, and amounts that were cleared.",
            "",
            "| Column | Missing before | Missing after |",
            "| --- | ---: | ---: |",
        ]
    )
    for column in EXPECTED_COLUMNS:
        before = stats["missing_before"][column]
        after = stats["missing_after"][column]
        lines.append(f"| {column} | {before} | {after} |")
    lines.append("")

    lines.extend(
        [
            "## Invalid values",
            "",
            "These cells could not be repaired, so the cleaned cell is empty.",
            "",
        ]
    )
    if stats["invalid"]:
        lines.append("| Source row | Column | Original value | Reason |")
        lines.append("| ---: | --- | --- | --- |")
        for item in stats["invalid"]:
            original = _markdown_cell(item["original"])
            lines.append(
                f"| {item['source_row']} | {item['column']} | {original} | {item['reason']} |"
            )
    else:
        lines.append("None.")
    lines.append("")

    lines.extend(
        [
            "## Cities to review",
            "",
            "These names are not in the city alias list. They were title-cased",
            "and kept. Add a spelling to CITY_ALIASES in cleaner.py to map them",
            "next time.",
            "",
        ]
    )
    if stats["unmapped_cities"]:
        lines.append("| Source row | Original value | Kept as |")
        lines.append("| ---: | --- | --- |")
        for item in stats["unmapped_cities"]:
            original = _markdown_cell(item["original"])
            lines.append(
                f"| {item['source_row']} | {original} | {item['kept_as']} |"
            )
    else:
        lines.append("None.")
    lines.append("")
    return "\n".join(lines)


def main(argv=None):
    args = parse_args(argv)
    input_path = Path(args.input_csv)
    output_dir = Path(args.output_dir)

    try:
        raw = load_input(input_path)
        cleaned, stats = clean_orders(raw)
    except (FileNotFoundError, UnicodeDecodeError, ValueError) as error:
        print(error, file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "cleaned.csv"
    excel_path = output_dir / "cleaned.xlsx"
    report_path = output_dir / "report.md"

    try:
        write_csv(cleaned, csv_path)
        write_excel(cleaned, excel_path)
        report_path.write_text(build_report(stats, input_path.as_posix()), encoding="utf-8")
    except PermissionError:
        print(
            f"Could not write into {output_dir}. Close the files if they are open and try again.",
            file=sys.stderr,
        )
        return 1

    print(
        f"Cleaned {stats['rows_before']} -> {stats['rows_after']} rows "
        f"({stats['duplicates_removed']} duplicates removed)."
    )
    print(f"Wrote {csv_path}")
    print(f"Wrote {excel_path}")
    print(f"Wrote {report_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
